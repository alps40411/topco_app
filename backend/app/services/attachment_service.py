# backend/app/services/attachment_service.py
# 附件下載與文字擷取服務。
# PDF / 圖片使用 Claude 視覺擷取，Office 檔與純文字檔使用本地解析（零 API 成本），
# 擷取結果供 Claude / Phison 兩條潤飾路徑共用。

import asyncio
import io
import logging
import os
import tempfile
from pathlib import Path
from typing import List

import aiofiles
import aiohttp

from app.core.config import settings
from app.services import claude_ai_service

logger = logging.getLogger(__name__)

# 擷取文字長度上限（避免 prompt 過長）
MAX_EXTRACT_CHARS = 10000

# Claude 單張圖片大小上限為 5MB，超過時先用 Pillow 壓縮
_MAX_IMAGE_BYTES = int(4.5 * 1024 * 1024)

_IMAGE_MEDIA_TYPES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".gif": "image/gif",
    ".webp": "image/webp",
}

# 需要先轉檔才能給 Claude 的圖片格式
_CONVERT_IMAGE_EXTS = {".bmp", ".tiff", ".tif"}

_TEXT_EXTS = {".txt", ".md", ".csv"}


async def extract_text_from_commonapi_url(url_path: str, file_name: str) -> str:
    """
    從 CommonAPI URL 下載檔案並提取文字內容

    Args:
        url_path: CommonAPI 檔案路徑或完整 URL
        file_name: 檔案名稱

    Returns:
        提取的文字內容，失敗時返回空字串
    """
    logger.info(f"從 CommonAPI 下載檔案: {file_name}")

    try:
        # 組合完整的 CommonAPI URL
        if url_path.startswith('/'):
            # 相對路徑：使用內網 API 基礎 URL
            full_url = f"{settings.COMMONAPI_BASE_URL}{url_path}"
        else:
            # 完整 URL：如果是外網域名，替換為內網 IP（避免認證問題）
            if 'portal.topco-global.com/tap1-98/CommonApi' in url_path:
                full_url = url_path.replace(
                    'https://portal.topco-global.com/tap1-98/CommonApi',
                    settings.COMMONAPI_BASE_URL
                )
                logger.info(f"替換外網域名為內網 IP: {settings.COMMONAPI_BASE_URL}")
            else:
                full_url = url_path

        logger.info(f"下載 URL: {full_url}")

        # 下載檔案到暫存目錄
        async with aiohttp.ClientSession() as session:
            async with session.get(full_url) as response:
                if response.status != 200:
                    logger.error(f"下載檔案失敗: HTTP {response.status}")
                    return ""

                # 創建暫存檔案
                file_ext = os.path.splitext(file_name)[1] or '.tmp'
                temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=file_ext)
                temp_path = temp_file.name

                # 讀取檔案內容
                content = await response.read()
                file_size = len(content)
                content_type = response.headers.get('Content-Type', 'unknown')

                # 驗證下載的檔案
                if file_ext.lower() == '.pdf' and content[:4] != b'%PDF':
                    logger.error(f"下載的檔案不是有效的 PDF！Content-Type: {content_type}")
                    logger.error(f"可能是認證失敗或重定向頁面，檔案開頭: {content[:100].decode('utf-8', errors='ignore')}")
                    return ""

                # 寫入暫存檔案
                temp_file.write(content)
                temp_file.close()
                logger.info(f"檔案已下載: {file_size} bytes, Content-Type: {content_type}")

        try:
            logger.info(f"開始提取檔案文字: {file_name}")
            extracted_text = await extract_text_from_file(temp_path)

            if extracted_text:
                logger.info(f"文字提取成功: {file_name}, 長度: {len(extracted_text)} 字符")
            else:
                logger.warning(f"文字提取失敗或檔案為空: {file_name}")

            return extracted_text
        finally:
            # 清理暫存檔案
            try:
                os.unlink(temp_path)
                logger.info(f"已清理暫存檔案: {temp_path}")
            except Exception as e:
                logger.warning(f"清理暫存檔案失敗: {e}")

    except Exception as e:
        import traceback
        logger.error(f"從 CommonAPI 提取檔案內容失敗: {file_name}")
        logger.error(f"錯誤詳情: {str(e)}")
        logger.debug(f"完整錯誤堆疊:\n{traceback.format_exc()}")
        return ""


async def extract_text_from_file(file_path: str) -> str:
    """
    從檔案中提取文字內容（限制 10000 字符），失敗時返回空字串。

    - .txt/.md/.csv：直接讀取（UTF-8 / Big5）
    - .docx/.xlsx/.pptx：本地解析
    - .pdf / 圖片：Claude 視覺擷取
    """
    logger.info(f"開始提取檔案內容: {file_path}")

    try:
        file_path_obj = Path(file_path)
        if not file_path_obj.exists():
            logger.error(f"檔案不存在: {file_path}")
            return ""

        file_ext = file_path_obj.suffix.lower()
        logger.info(f"檔案類型: {file_ext}")

        if file_ext in _TEXT_EXTS:
            return await _read_plain_text(file_path)

        if file_ext == ".docx":
            return (await asyncio.to_thread(_extract_docx, file_path))[:MAX_EXTRACT_CHARS]

        if file_ext == ".xlsx":
            return (await asyncio.to_thread(_extract_xlsx, file_path))[:MAX_EXTRACT_CHARS]

        if file_ext == ".pptx":
            return (await asyncio.to_thread(_extract_pptx, file_path))[:MAX_EXTRACT_CHARS]

        if file_ext == ".pdf":
            async with aiofiles.open(file_path, "rb") as f:
                data = await f.read()
            text = await claude_ai_service.extract_text_from_document(
                data, "application/pdf", file_path_obj.name
            )
            return text[:MAX_EXTRACT_CHARS]

        if file_ext in _IMAGE_MEDIA_TYPES or file_ext in _CONVERT_IMAGE_EXTS:
            async with aiofiles.open(file_path, "rb") as f:
                data = await f.read()
            media_type = _IMAGE_MEDIA_TYPES.get(file_ext, "")
            # bmp/tiff 或超過大小上限時，轉為 JPEG 再送
            if file_ext in _CONVERT_IMAGE_EXTS or len(data) > _MAX_IMAGE_BYTES:
                data = await asyncio.to_thread(_normalize_image, data)
                media_type = "image/jpeg"
            text = await claude_ai_service.extract_text_from_document(
                data, media_type, file_path_obj.name
            )
            return text[:MAX_EXTRACT_CHARS]

        logger.warning(f"不支援的檔案類型: {file_ext}")
        return ""

    except Exception as e:
        logger.error(f"檔案內容提取失敗 {file_path}: {str(e)}")
        return ""


async def _read_plain_text(file_path: str) -> str:
    try:
        logger.info("嘗試以UTF-8讀取純文字檔案")
        async with aiofiles.open(file_path, 'r', encoding='utf-8') as f:
            content = await f.read()
            logger.info(f"成功讀取純文字檔案: {len(content)} 字符")
            return content[:MAX_EXTRACT_CHARS]
    except UnicodeDecodeError:
        try:
            logger.info("UTF-8失敗，嘗試以Big5讀取")
            async with aiofiles.open(file_path, 'r', encoding='big5') as f:
                content = await f.read()
                logger.info(f"以Big5成功讀取純文字檔案: {len(content)} 字符")
                return content[:MAX_EXTRACT_CHARS]
        except Exception as e:
            logger.error(f"Big5讀取也失敗: {str(e)}")
            return ""
    except Exception as e:
        logger.error(f"讀取純文字檔案失敗: {str(e)}")
        return ""


def _extract_docx(file_path: str) -> str:
    from docx import Document

    doc = Document(file_path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_xlsx(file_path: str) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    parts = []
    try:
        for sheet in wb.worksheets:
            parts.append(f"【工作表：{sheet.title}】")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if cells:
                    parts.append(" | ".join(cells))
                if sum(len(p) for p in parts) > MAX_EXTRACT_CHARS:
                    return "\n".join(parts)
    finally:
        wb.close()
    return "\n".join(parts)


def _extract_pptx(file_path: str) -> str:
    from pptx import Presentation

    prs = Presentation(file_path)
    parts = []
    for idx, slide in enumerate(prs.slides, start=1):
        parts.append(f"【投影片 {idx}】")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
    return "\n".join(parts)


def _normalize_image(data: bytes) -> bytes:
    """將圖片轉為 JPEG 並限制尺寸，確保符合 Claude 的圖片限制。"""
    from PIL import Image

    with Image.open(io.BytesIO(data)) as img:
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        img.thumbnail((3000, 3000))
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=85)
        return buf.getvalue()


async def process_attachments_for_ai(attachment_records: List[dict]) -> List[str]:
    """
    處理附件列表，提取標記為 AI 參考的檔案內容

    Args:
        attachment_records: 附件記錄列表，每個記錄應包含：
            - file_name: 檔案名稱
            - file_path: CommonAPI 檔案路徑或 URL
            - is_selected_for_ai: 是否標記為 AI 參考

    Returns:
        提取的文字內容列表（包含檔案來源標識）
    """
    reference_texts = []
    logger.info(f"開始處理附件，總數: {len(attachment_records)}")

    for attachment in attachment_records:
        # 只處理標記為 AI 參考的附件
        if not attachment.get('is_selected_for_ai', False):
            logger.debug(f"跳過未標記為AI參考的附件: {attachment.get('file_name', '未知')}")
            continue

        file_path = attachment.get('file_path')
        file_name = attachment.get('file_name', '未知檔案')

        if not file_path:
            logger.warning(f"附件 {file_name} 缺少檔案路徑，跳過")
            continue

        logger.info(f"處理 AI 參考檔案: {file_name}")

        try:
            # 從 CommonAPI 下載並提取檔案內容
            extracted_text = await extract_text_from_commonapi_url(file_path, file_name)

            if extracted_text:
                # 添加檔案來源標識，方便 AI 理解內容來源
                formatted_text = f"【檔案：{file_name}】\n{extracted_text}"
                reference_texts.append(formatted_text)
                logger.info(f"✅ 成功提取: {file_name} ({len(extracted_text)} 字符)")
            else:
                logger.warning(f"⚠️ 提取失敗或檔案為空: {file_name}")

        except Exception as e:
            logger.error(f"處理檔案 {file_name} 時發生錯誤: {str(e)}")
            # 繼續處理其他檔案，不中斷整個流程
            continue

    logger.info(f"✅ 附件處理完成，成功提取 {len(reference_texts)}/{len(attachment_records)} 個檔案")
    return reference_texts
